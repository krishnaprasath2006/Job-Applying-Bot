"""Build the synthetic PDF and DOCX resume fixtures.

Run from the project root with the project virtualenv::

    python tests/fixtures/make_binary_fixtures.py

Both files are generated rather than committed as opaque binaries so that what
is inside them is reviewable in a diff. The generated files are test fixtures
and must never contain real personal data.

The PDF is written with a small hand-rolled PDF 1.4 writer because no PDF
library is a runtime dependency of this project: adding reportlab purely to
produce a test fixture would make a build-time dependency out of a test-only
concern. pdfplumber reads the result back, which is verified immediately after
writing.
"""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
if str(ROOT / "src") not in sys.path:
    sys.path.insert(0, str(ROOT / "src"))

SYNTHETIC_MARKER = "SYNTHETIC TEST DATA ONLY"

RESUME_LINES = [
    SYNTHETIC_MARKER,
    "",
    "Jane Synthetic",
    "Machine Learning Engineer (synthetic candidate)",
    "jane.synthetic@example.invalid | +91 00000 00000",
    "Remote, India | github.com/invalid-synthetic",
    "",
    "SUMMARY",
    "Synthetic candidate with placeholder experience used only in automated",
    "tests. Not a real person, employer, or application.",
    "",
    "SKILLS",
    "Python, SQL, Pydantic, SQLite, RAG, vector search, pdfplumber, pytest",
    "",
    "EXPERIENCE",
    "Machine Learning Engineer, Northwind Analytics (synthetic)",
    "2023-01 to present",
    "- Built a retrieval-augmented generation pipeline over synthetic documents.",
    "- Designed evaluation harnesses comparing deterministic parsers to model output.",
    "",
    "Data Engineer, Contoso Synthetic Labs (synthetic)",
    "2020-06 to 2022-12",
    "- Owned document ingestion for PDF and DOCX sources.",
    "",
    "EDUCATION",
    "B.E. Computer Science, Example Institute (synthetic), 2020",
    "",
    "CERTIFICATIONS",
    "Synthetic Certificate in Testing Discipline (synthetic)",
    "",
    "PROJECTS",
    "synthetic-rag-demo - a placeholder repository used in tests.",
    "",
    "END OF " + SYNTHETIC_MARKER,
]


def build_pdf(path: Path, lines: list[str]) -> None:
    """Write a minimal, valid single-page PDF containing ``lines``."""
    def escape(text: str) -> str:
        # Backslash, parens must be escaped for a PDF literal string.
        return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")

    content_parts = ["BT", "/F1 11 Tf", "14 TL", "72 740 Td"]
    for line in lines:
        content_parts.append(f"({escape(line)}) Tj")
        content_parts.append("T*")
    content_parts.append("ET")
    content = "\n".join(content_parts).encode("latin-1", "replace")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"

    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode()

    path.write_bytes(bytes(out))


def build_docx(path: Path, lines: list[str]) -> None:
    """Write a DOCX containing ``lines``, including a table so the table path is covered.

    Timestamps are pinned so the output is byte-for-byte reproducible. python-docx
    otherwise stamps the current time into the package metadata, which would make
    every regeneration differ and turn a real fixture diff into noise.
    """
    from docx import Document

    document = Document()
    for line in lines:
        document.add_paragraph(line)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Synthetic Skill"
    table.cell(0, 1).text = "Synthetic Level"
    table.cell(1, 0).text = "Python"
    table.cell(1, 1).text = "SYNTHETIC"

    fixed = datetime(2026, 1, 1, 0, 0, 0)
    properties = document.core_properties
    properties.created = fixed
    properties.modified = fixed
    properties.last_printed = fixed
    properties.revision = 1

    document.save(str(path))
    _pin_zip_timestamps(path)


def _pin_zip_timestamps(path: Path) -> None:
    """Rewrite a DOCX zip with fixed entry timestamps.

    A DOCX is a zip, and zipfile stamps each entry with the current clock, so
    two runs produce different bytes even when the content is identical. Fixing
    the timestamps is what makes the regeneration test meaningful.
    """
    import zipfile

    fixed = (2026, 1, 1, 0, 0, 0)
    with zipfile.ZipFile(path, "r") as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for info, data in entries:
            pinned = zipfile.ZipInfo(info.filename, date_time=fixed)
            pinned.compress_type = info.compress_type
            pinned.external_attr = info.external_attr
            pinned.create_system = info.create_system
            target.writestr(pinned, data)


def main() -> int:
    """Write both fixtures and verify them.

    An output directory may be passed as the single argument. It defaults to
    this file's own directory. The argument exists so the regeneration test can
    run the generator from a temporary directory and compare the result with the
    committed fixtures, which is only meaningful if the output is byte-for-byte
    reproducible.
    """
    out_dir = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else HERE
    out_dir.mkdir(parents=True, exist_ok=True)
    pdf_path = out_dir / "synthetic_resume.pdf"
    docx_path = out_dir / "synthetic_resume.docx"

    build_pdf(pdf_path, RESUME_LINES)
    build_docx(docx_path, RESUME_LINES)
    print(f"wrote {pdf_path} ({pdf_path.stat().st_size} bytes)")
    print(f"wrote {docx_path} ({docx_path.stat().st_size} bytes)")

    # Verify the PDF is readable by the same library the parser uses.
    import pdfplumber

    with pdfplumber.open(str(pdf_path)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert SYNTHETIC_MARKER in text, "PDF did not round-trip through pdfplumber"
    print("PDF verified via pdfplumber:", len(text), "characters extracted")

    from resumes.parser import extract_text

    assert SYNTHETIC_MARKER in extract_text(pdf_path), "parser failed on generated PDF"
    assert SYNTHETIC_MARKER in extract_text(docx_path), "parser failed on generated DOCX"
    print("Both generated fixtures parse successfully.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
