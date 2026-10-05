"""Text extraction and deterministic section parsing.

Four formats, one interface. Every extractor is allowed to fail, and none of
them invents text: a PDF that yields nothing raises
:class:`~core.errors.ResumeParseError` rather than returning an empty string
that a later stage would treat as a successful parse.

Section detection is keyword-based and deterministic. It does not use an LLM
and it does not guess: an unrecognised document produces no sections and is
reported as such, so a human can see that parsing failed instead of the
assistant quietly inventing structure.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Callable, Optional

from core.enums import ResumeFileType, ResumeStatus, SectionType
from core.errors import ResumeParseError
from core.hashing import utc_now
from core.logging_config import get_logger
from resumes.models import PARSER_NAME, PARSER_VERSION, Resume, ResumeSection
from resumes.hashing import MIN_TEXT_CHARS, ensure_ingestable, hash_file, hash_text

# Structured logger. Defined because _page_count() reports its own best-effort
# failures at debug level, and an undefined name there would raise NameError in
# the middle of an otherwise successful ingest -- turning "the page count is
# unavailable" into "this resume cannot be parsed".
log = get_logger("resumes.parser")

__all__ = [
    "ResumeParser",
    "extract_text",
    "extract_txt",
    "extract_md",
    "extract_pdf",
    "extract_docx",
    "PARSER_NAME",
    "PARSER_VERSION",
]

_ENCODINGS = ("utf-8", "utf-8-sig", "cp1252", "latin-1")


# --------------------------------------------------------------------------
# Per-format extractors
# --------------------------------------------------------------------------
def _require_minimum(text: str, path: Path, kind: str) -> str:
    """Refuse text too short to be a resume.

    Applied by every extractor, not just the binary ones. A five-character text
    file is not a resume with no sections, it is a mistake, and storing it
    would quietly produce an empty record that looks ingested.
    """
    if len(text) < MIN_TEXT_CHARS:
        raise ResumeParseError(
            f"{kind} file has too little text to be a resume",
            path=str(path),
            char_count=len(text),
            minimum=MIN_TEXT_CHARS,
        )
    return text


def extract_txt(path: Path) -> str:
    """Read a plain-text file, tolerating common legacy encodings.

    Raises:
        ResumeParseError: If the file cannot be decoded, is empty, or holds too
            little text to be a resume.
    """
    raw = _read_bytes(path, "txt")
    for encoding in _ENCODINGS:
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError:
            continue
        text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
        if text:
            return _require_minimum(text, path, "text")
    raise ResumeParseError("could not decode text file", path=str(path))


def extract_md(path: Path) -> str:
    """Read a Markdown file.

    The markdown syntax is kept in the preserved text so a human reading the
    stored copy sees the original document, not a lossy conversion.
    """
    return _require_minimum(extract_txt(path), path, "markdown")


def extract_pdf(path: Path) -> str:
    """Extract text from a PDF using pdfplumber.

    Raises:
        ResumeParseError: If pdfplumber is missing, the PDF is unreadable, or
            the pages contain no extractable text. Scanned-image PDFs land
            here: they need OCR, which this phase does not do, and failing
            loudly is correct because a silent empty parse would look like a
            resume with nothing in it.
    """
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ResumeParseError(
            "pdfplumber is not installed; install it to ingest PDFs",
            path=str(path),
        ) from exc

    pages: list[str] = []
    try:
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
    except Exception as exc:
        raise ResumeParseError(
            "could not read PDF",
            path=str(path),
            cause=type(exc).__name__,
        ) from exc

    text = "\n\n".join(p.strip() for p in pages if p and p.strip()).strip()
    if len(text) < MIN_TEXT_CHARS:
        raise ResumeParseError(
            "PDF contains no extractable text; it is probably a scanned image "
            "and needs OCR before it can be ingested",
            path=str(path),
            char_count=len(text),
            minimum=MIN_TEXT_CHARS,
        )
    return text


def extract_docx(path: Path) -> str:
    """Extract text from a .docx using python-docx.

    Paragraphs and table cells are both included, since many resumes put
    employment history into a table.

    Raises:
        ResumeParseError: If python-docx is missing, the file is not a valid
            docx, or no text could be read.
    """
    try:
        import docx  # type: ignore[import-untyped]
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ResumeParseError(
            "python-docx is not installed; install it to ingest DOCX files",
            path=str(path),
        ) from exc

    parts: list[str] = []
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ResumeParseError(
            "could not open DOCX; the file may not be a real .docx",
            path=str(path),
            cause=type(exc).__name__,
        ) from exc

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(parts).strip()
    if len(text) < MIN_TEXT_CHARS:
        raise ResumeParseError(
            "DOCX contains too little text to be a resume",
            path=str(path),
            char_count=len(text),
            minimum=MIN_TEXT_CHARS,
        )
    return text


_EXTRACTORS: dict[ResumeFileType, Callable[[Path], str]] = {
    ResumeFileType.TXT: extract_txt,
    ResumeFileType.MD: extract_md,
    ResumeFileType.PDF: extract_pdf,
    ResumeFileType.DOCX: extract_docx,
}


def _read_bytes(path: Path, kind: str) -> bytes:
    try:
        return Path(path).read_bytes()
    except OSError as exc:
        raise ResumeParseError(
            f"could not read {kind} file", path=str(path), cause=str(exc)
        ) from exc


def extract_text(path: Path | str, file_type: Optional[ResumeFileType] = None) -> str:
    """Extract text from any supported resume file.

    Args:
        path: File to read.
        file_type: Detected type; inferred from the suffix when omitted.

    Returns:
        The extracted text.

    Raises:
        ResumeParseError: If extraction fails for any reason.
    """
    target, detected, _size = ensure_ingestable(path)
    kind = file_type or detected
    extractor = _EXTRACTORS.get(kind)
    if extractor is None:
        raise ResumeParseError("no extractor for this file type", path=str(target), file_type=kind.value)
    return extractor(target)


# --------------------------------------------------------------------------
# Deterministic section parsing
# --------------------------------------------------------------------------
_HEADING_RE = re.compile(
    r"^\s{0,6}(?:#{1,6}\s*)?(?P<label>[A-Z][A-Za-z &/]{2,40})\s*:?\s*$",
    re.IGNORECASE,
)
_BULLET_RE = re.compile(r"^\s*(?:[-*+\u2022\u2023\u25aa\u2013]|\d+[.)])\s+")
_DATE_RANGE_RE = re.compile(
    r"(?P<start>(?:19|20)\d{2}(?:[-/](?:0[1-9]|1[0-2]))?)"
    r"\s*(?:-|–|—|to|until)\s*"
    r"(?P<end>(?:19|20)\d{2}(?:[-/](?:0[1-9]|1[0-2]))?|present|current|now)",
    re.IGNORECASE,
)

_SECTION_RULES: tuple[tuple[SectionType, tuple[str, ...]], ...] = (
    (SectionType.SUMMARY, ("summary", "profile", "objective", "about")),
    (SectionType.EDUCATION, ("education", "academic qualification", "qualifications", "degree")),
    (
        SectionType.EXPERIENCE,
        ("experience", "employment", "work history", "professional experience", "internship"),
    ),
    (SectionType.SKILLS, ("skills", "technical skills", "technologies", "competencies", "tools")),
    (SectionType.PROJECTS, ("projects", "personal projects")),
    (SectionType.CERTIFICATIONS, ("certifications", "certificates", "licenses", "credentials")),
    (SectionType.ACHIEVEMENTS, ("achievements", "awards", "honors", "publications")),
    (SectionType.LINKS, ("links", "references", "online profiles")),
    (SectionType.CONTACT, ("contact", "contact details")),
)


def _classify_heading(label: str) -> Optional[SectionType]:
    lowered = label.strip().lower().rstrip(":")
    for section_type, keywords in _SECTION_RULES:
        for keyword in keywords:
            if lowered == keyword or lowered.startswith(keyword):
                return section_type
    return None


def _split_items(block: str) -> list[dict[str, Any]]:
    """Split a section body into items.

    Each item is either a bullet line or a paragraph separated by a blank
    line. Only structure that is actually present is returned; nothing is
    synthesised.
    """
    items: list[dict[str, Any]] = []
    current: list[str] = []

    def flush() -> None:
        if not current:
            return
        joined = "\n".join(current).strip()
        if joined:
            items.append(_describe_item(joined))
        current.clear()

    for line in block.splitlines():
        if not line.strip():
            flush()
            continue
        if _BULLET_RE.match(line):
            flush()
            items.append(_describe_item(_BULLET_RE.sub("", line).strip()))
            continue
        current.append(line)
    flush()
    return items


def _describe_item(text: str) -> dict[str, Any]:
    """Describe one item using only what is literally present.

    ``title`` is taken from the first fragment before a separator, and
    ``date_range`` only when a date pattern is actually found. Absent
    information is left absent.
    """
    match = _DATE_RANGE_RE.search(text)
    date_range: Optional[dict[str, str]] = None
    if match:
        date_range = {
            "start": match.group("start"),
            "end": match.group("end").lower(),
        }

    head = text.split("|")[0].split(" at ")[0].strip()
    pieces = [p.strip() for p in re.split(r"\s{2,}|\s+[|·]\s+|\s+-\s+", head) if p.strip()]
    return {
        "text": text,
        "title": pieces[0] if pieces else text[:80],
        "date_range": date_range,
    }


class ResumeParser:
    """Turns a resume file into a :class:`~resumes.models.Resume`.

    Usage::

        parser = ResumeParser(candidate_id="primary")
        resume = parser.parse("C:/resumes/ai_engineer.pdf", variant="AI Engineer")

    The original file is only read. Nothing is moved, renamed, or overwritten.
    """

    version = PARSER_VERSION

    def __init__(self, candidate_id: str = "primary", *, is_synthetic: bool = False) -> None:
        self.candidate_id = candidate_id
        self.is_synthetic = is_synthetic

    def parse(
        self,
        path: Path | str,
        *,
        variant: Optional[str] = None,
        text: Optional[str] = None,
    ) -> Resume:
        """Parse a resume file.

        Args:
            path: The resume file. Read only.
            variant: Optional intended use, e.g. ``AI Engineer``. Stored, not
                acted on.
            text: Pre-extracted text, used by tests and by callers that already
                have the content. When supplied, extraction is skipped but
                hashing still runs against the real file.

        Returns:
            A :class:`Resume` with ``status=PARSED`` and sections populated.

        Raises:
            ResumeValidationError: If the file is missing, unsupported, or empty.
            ResumeParseError: If text extraction fails.
        """
        target, file_type, size = ensure_ingestable(path)
        file_hash = hash_file(target)
        started = utc_now()

        extracted = text if text is not None else extract_text(target, file_type)

        resume = Resume(
            candidate_id=self.candidate_id,
            variant=variant,
            filename=target.name,
            file_type=file_type,
            file_hash=file_hash,
            file_size_bytes=size,
            char_count=len(extracted),
            content_hash=hash_text(extracted),
            storage_path=str(target.resolve()),
            raw_text=extracted,
            status=ResumeStatus.PARSED,
            parser_name=PARSER_NAME,
            parser_version=self.version,
            is_synthetic=self.is_synthetic,
            page_count=self._page_count(target, file_type),
            metadata={
                "parsed_at": utc_now().isoformat(),
                "parse_duration_ms": int((utc_now() - started).total_seconds() * 1000),
                "section_parser": "deterministic_keywords",
                "uses_llm": False,
            },
        )
        resume.sections = self.parse_sections(extracted, resume_id=resume.id)
        return resume

    def parse_sections(self, text: str, *, resume_id: str = "") -> list[ResumeSection]:
        """Split text into sections by heading keyword.

        Returns:
            Sections in document order. Possibly empty when no heading matches,
            which is a truthful result rather than a failure to report.
        """
        lines = text.splitlines()
        boundaries: list[tuple[int, SectionType, str]] = []

        for index, line in enumerate(lines):
            match = _HEADING_RE.match(line)
            if not match:
                continue
            label = match.group("label")
            section_type = _classify_heading(label)
            if section_type is not None:
                boundaries.append((index, section_type, label.strip().rstrip(":")))

        sections: list[ResumeSection] = []
        for position, (start, section_type, label) in enumerate(boundaries):
            end = boundaries[position + 1][0] if position + 1 < len(boundaries) else len(lines)
            body = "\n".join(lines[start + 1 : end]).strip()
            sections.append(
                ResumeSection(
                    resume_id=resume_id,
                    section_type=section_type,
                    position=position,
                    heading=label,
                    raw_text=body,
                    items=_split_items(body),
                )
            )
        return sections

    def _page_count(self, path: Path, file_type: ResumeFileType) -> Optional[int]:
        """Best-effort PDF page count.

        Page count is optional metadata. The resume text has already been
        extracted successfully before this runs, so a failure here means only
        the count is unavailable: return ``None``, record why at debug level,
        and never abort ingestion. ``PDFSyntaxError`` and ``PSException`` are
        pdfminer/pdfplumber's own parse errors; ``OSError`` and ``ValueError``
        cover encrypted, corrupt, or non-file inputs.
        """
        if file_type is not ResumeFileType.PDF:
            return None
        try:
            import pdfplumber
            from pdfminer.psparser import PSException
            from pdfminer.pdfparser import PDFSyntaxError
        except ImportError as exc:  # pragma: no cover - pdfplumber checked at extract time
            log.debug("pdfplumber unavailable for page count", extra={"error": str(exc)})
            return None
        try:
            with pdfplumber.open(str(path)) as pdf:
                return len(pdf.pages)
        except (OSError, ValueError, PDFSyntaxError, PSException) as exc:
            log.debug(
                "could not determine PDF page count",
                extra={"resume": {"path": str(path), "error": type(exc).__name__}},
            )
            return None
